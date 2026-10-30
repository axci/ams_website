"""Turn client order spreadsheets (received by email) into site Orders.

Different clients send different layouts, so the importer is driven by a list of
*source profiles* (``settings.EMAIL_ORDER_IMPORT["SOURCES"]``). Each profile names
an IMAP folder, a file format, how the sender maps to a site account, and which
product field to match on. Two formats are supported today:

``nkoil_sections``
    An .xlsx where each «Точка продаж: …» row starts a section (one Order each);
    item lines carry the client's own code (matched to ``Product.nkoil_code``)
    and a «К заказу» quantity. Columns are located by their header labels, since
    the client ships 4- and 9-column variants.

``price_list``
    A filled catalog price list (.xls or .xlsx): the whole file is one order, and
    every row with a positive «Заказ» becomes a line whose «Код» is our
    ``Product.sku``.

Parsing is pure and unit-testable; :func:`create_orders_for_email` does the DB
work. IMAP fetching lives in the ``import_email_orders`` management command.
"""

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal

import openpyxl
from django.conf import settings
from django.db import transaction

# ---- nkoil_sections format -------------------------------------------------- #
POINT_PREFIX = "Точка продаж:"
# Columns are located by these header labels, because the client's layout varies
# — one file is 4 columns with «К заказу» in column D, another 9 columns with a
# «Цена (закупка)» column that pushes «К заказу» to column E.
CODE_HEADER = "Код товара"
NAME_HEADER = "Наименование товара"
QTY_HEADER = "К заказу"
DEFAULT_COLS = {"code": 0, "name": 1, "qty": 3}
# A trailing position count the client appends to a point-of-sale name, e.g.
# "АЗС №25 … (10)" — the "(10)" is the number of lines, not part of the name.
_TRAILING_COUNT = re.compile(r"\s*\(\d+\)\s*$")

# ---- price_list format ------------------------------------------------------ #
PL_CODE_HEADER = "Код"      # our Product.sku
PL_QTY_HEADER = "Заказ"     # ordered quantity
PL_NAME_HEADER = "Наименование"


# --------------------------------------------------------------------------- #
# Parsing (pure — no database, no IMAP)
# --------------------------------------------------------------------------- #
@dataclass
class ParsedLine:
    code: str
    name: str
    quantity: int
    row: int  # 1-based source row, for error messages


@dataclass
class ParsedOrder:
    point_of_sale: str
    items: list = field(default_factory=list)


def _clean(value):
    """Normalise a cell to a trimmed string (numbers lose a trailing .0)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_qty(value):
    """Parse a quantity cell to a positive int, or None if not usable."""
    text = _clean(value).replace(",", ".")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    rounded = int(round(number))
    return rounded if rounded > 0 else None


def _header_columns(cells):
    """Map code/name/qty to 0-based indices from an nkoil_sections header row,
    falling back to the default layout for any label that isn't found."""
    labels = {i: _clean(c).lower() for i, c in enumerate(cells)}
    cols = dict(DEFAULT_COLS)
    for key, header in (("code", CODE_HEADER), ("name", NAME_HEADER), ("qty", QTY_HEADER)):
        for i, text in labels.items():
            if text == header.lower():
                cols[key] = i
                break
    return cols


def parse_sections(rows):
    """nkoil_sections: «Точка продаж» sections → one :class:`ParsedOrder` each.

    Columns are located from the header row; the code column (col A) anchors
    section markers and subtotal rows. Sections with no item lines are dropped.
    """
    orders = []
    current = None
    cols = dict(DEFAULT_COLS)

    def at(cells, key):
        i = cols[key]
        return cells[i] if 0 <= i < len(cells) else None

    for idx, row in enumerate(rows, start=1):
        cells = list(row) if row is not None else []
        code = _clean(at(cells, "code"))

        if code == CODE_HEADER:
            cols = _header_columns(cells)  # lock column positions from the header
            continue
        if code.startswith(POINT_PREFIX):
            name = _TRAILING_COUNT.sub("", code[len(POINT_PREFIX):].strip()).strip()
            current = ParsedOrder(point_of_sale=name)
            orders.append(current)
            continue
        if not code:
            continue  # subtotal / grand-total row (code column empty)

        if current is None:  # tolerate items before any marker row
            current = ParsedOrder(point_of_sale="")
            orders.append(current)
        current.items.append(
            ParsedLine(
                code=code,
                name=_clean(at(cells, "name")),
                quantity=_parse_qty(at(cells, "qty")) or 0,
                row=idx,
            )
        )

    return [order for order in orders if order.items]


def parse_price_list(rows):
    """price_list: a filled catalog price list → a single :class:`ParsedOrder`.

    Finds the header row (the one carrying «Код» and «Заказ»), then every row
    below with a positive «Заказ» becomes a line keyed by «Код» (our sku).
    """
    code_i = qty_i = name_i = None
    items = []
    for idx, row in enumerate(rows, start=1):
        cells = list(row) if row is not None else []
        if code_i is None:
            texts = [_clean(c) for c in cells]
            if PL_CODE_HEADER in texts and PL_QTY_HEADER in texts:
                code_i = texts.index(PL_CODE_HEADER)
                qty_i = texts.index(PL_QTY_HEADER)
                name_i = texts.index(PL_NAME_HEADER) if PL_NAME_HEADER in texts else None
            continue
        code = _clean(cells[code_i]) if code_i < len(cells) else ""
        quantity = _parse_qty(cells[qty_i]) if qty_i < len(cells) else None
        if not code or not quantity:
            continue  # category header / un-ordered row
        name = _clean(cells[name_i]) if (name_i is not None and name_i < len(cells)) else ""
        items.append(ParsedLine(code=code, name=name, quantity=quantity, row=idx))

    return [ParsedOrder(point_of_sale="", items=items)] if items else []


FORMATS = {
    "nkoil_sections": parse_sections,
    "price_list": parse_price_list,
}

# Back-compat alias (older tests import parse_order_rows).
parse_order_rows = parse_sections


def read_rows(filename, data):
    """First sheet of an .xlsx/.xls attachment as a list of row value-lists."""
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        import xlrd  # only needed for the legacy binary format

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(r) for r in range(sheet.nrows)]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_attachment(filename, data, fmt="nkoil_sections"):
    """Parse an attachment's bytes with the parser for ``fmt``."""
    return FORMATS.get(fmt, parse_sections)(read_rows(filename, data))


def parse_order_workbook(path_or_file, fmt="nkoil_sections"):
    """Parse an .xlsx path/file-like (used by tests) with the parser for ``fmt``."""
    wb = openpyxl.load_workbook(path_or_file, data_only=True, read_only=True)
    try:
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    return FORMATS.get(fmt, parse_sections)(rows)


# --------------------------------------------------------------------------- #
# Configuration / resolution helpers
# --------------------------------------------------------------------------- #
def get_config():
    """The EMAIL_ORDER_IMPORT settings dict (empty dict if unset)."""
    return getattr(settings, "EMAIL_ORDER_IMPORT", {}) or {}


def get_sources(config=None):
    """The list of source profiles, or a single legacy profile synthesised from
    the old flat keys if no SOURCES are configured."""
    config = config or get_config()
    sources = config.get("SOURCES")
    if sources:
        return sources
    return [{
        "name": "default",
        "folder": config.get("MAILBOX", "INBOX"),
        "format": "nkoil_sections",
        "match_by": "nkoil_code",
        "senders": config.get("ALLOWED_SENDERS", []),
        "sender_user_map": config.get("SENDER_USER_MAP", ""),
    }]


def _parse_sender_map(raw):
    """Parse "a@b=login1,c@d=login2" into {email: login} (lower-cased emails)."""
    mapping = {}
    for pair in (raw or "").split(","):
        pair = pair.strip()
        if "=" not in pair:
            continue
        email_addr, _, login = pair.partition("=")
        email_addr, login = email_addr.strip().lower(), login.strip()
        if email_addr and login:
            mapping[email_addr] = login
    return mapping


def sender_allowed(sender, source):
    """Whether a message from ``sender`` should be processed for this source.

    A source with an explicit ``senders`` whitelist accepts only those; a source
    without one (e.g. resolve-by-email) accepts any sender and relies on user
    resolution as the gate.
    """
    senders = source.get("senders")
    if not senders:
        return True
    return (sender or "").strip().lower() in {s.strip().lower() for s in senders}


def resolve_user(sender, source):
    """Find the site account an email's orders belong to, per the source profile.

    ``resolve_by_email`` → the account whose email equals the sender; otherwise
    the explicit sender→login map. Returns None when nothing matches, so the mail
    is reported and left for a human rather than guessed at.
    """
    from accounts.models import User

    sender = (sender or "").strip().lower()
    if not sender:
        return None
    if source.get("resolve_by_email"):
        return User.objects.filter(email__iexact=sender).first()
    smap = source.get("sender_user_map") or {}
    if isinstance(smap, str):
        smap = _parse_sender_map(smap)
    else:
        smap = {str(k).strip().lower(): v for k, v in smap.items()}
    login = smap.get(sender)
    return User.objects.filter(username=login).first() if login else None


def resolve_warehouse(config=None):
    """The default warehouse new imported orders are placed on."""
    from warehouses.models import Warehouse

    config = config or get_config()
    name = (config.get("DEFAULT_WAREHOUSE") or "").strip()
    if not name:
        return None
    return Warehouse.objects.filter(name=name).first()


def match_product(code, match_by="nkoil_code"):
    """Find a catalog product by a client code. ``match_by`` is the Product field
    to match on: 'nkoil_code' (the client's own code) or 'sku' (our code)."""
    from catalog.models import Product

    code = (code or "").strip()
    if not code:
        return None
    field_name = match_by if match_by in ("nkoil_code", "sku") else "nkoil_code"
    qs = Product.objects.filter(**{field_name: code})
    return qs.filter(is_active=True).first() or qs.first()


# --------------------------------------------------------------------------- #
# Order creation
# --------------------------------------------------------------------------- #
@dataclass
class ImportResult:
    created_orders: list = field(default_factory=list)   # Order pks
    sections: int = 0
    lines: int = 0
    matched: int = 0
    unmatched: list = field(default_factory=list)        # (code, name, qty, row)
    skipped: list = field(default_factory=list)          # bad/zero qty lines
    notes: list = field(default_factory=list)

    def summary(self):
        return (
            f"секций: {self.sections}, заказов создано: {len(self.created_orders)}, "
            f"строк: {self.lines}, сопоставлено: {self.matched}, "
            f"не найдено: {len(self.unmatched)}, пропущено: {len(self.skipped)}"
        )

    def render_report(self):
        """A human-readable report stored on the ImportedOrderEmail row."""
        lines = [self.summary()]
        pks = [pk for pk in self.created_orders if pk is not None]
        if pks:
            lines.append("Заказы: " + ", ".join(f"№{pk}" for pk in pks))
        for code, name, qty, row in self.unmatched:
            lines.append(f"НЕ НАЙДЕН товар [стр.{row}]: {code} ×{qty} — {name}")
        for code, name, qty, row in self.skipped:
            lines.append(f"ПРОПУЩЕНО (нет кол-ва) [стр.{row}]: {code} — {name}")
        lines.extend(self.notes)
        return "\n".join(lines)


def _order_comment(source_label, point_of_sale):
    parts = ["Импортирован из письма"]
    if source_label:
        parts[0] += f" ({source_label})"
    if point_of_sale:
        parts.append(f"Точка продаж: {point_of_sale}")
    return "\n".join(parts)


def create_orders_for_email(
    sections, *, user, warehouse, match_by="nkoil_code",
    source_email=None, source_label="", dry_run=False,
):
    """Create one pending Order per section for ``user`` on ``warehouse``.

    Lines are matched to products by ``match_by`` ('nkoil_code' or 'sku');
    unmatched or quantity-less lines are recorded in the result rather than
    failing the import. A section that yields no matched line creates no order.
    With ``dry_run`` nothing is written. Everything is wrapped in one transaction.
    """
    from .models import Order, OrderItem

    result = ImportResult(sections=len(sections))

    # A B2B buyer usually has exactly one company — attach it so the 1С export
    # resolves the контрагент; leave it unset when ambiguous.
    companies = list(user.companies.all()[:2])
    company = companies[0] if len(companies) == 1 else None

    with transaction.atomic():
        for section in sections:
            matched_lines = []
            for line in section.items:
                result.lines += 1
                if line.quantity <= 0:
                    result.skipped.append((line.code, line.name, line.quantity, line.row))
                    continue
                product = match_product(line.code, match_by)
                if product is None:
                    result.unmatched.append((line.code, line.name, line.quantity, line.row))
                    continue
                price = product.price_for_user(user)
                if price is None:
                    price = Decimal("0")
                    result.notes.append(
                        f"нет цены: {product.sku} ({line.code}) — цена 0"
                    )
                matched_lines.append((product, price, line.quantity))
                result.matched += 1

            if not matched_lines:
                continue  # don't create an empty order

            if dry_run:
                result.created_orders.append(None)  # a would-be order, no DB write
                continue

            order = Order.objects.create(
                user=user,
                warehouse=warehouse,
                company=company,
                status=Order.Status.PENDING,
                comment=_order_comment(source_label, section.point_of_sale),
                imported_from_email=source_email,
            )
            for product, price, quantity in matched_lines:
                OrderItem.objects.create(
                    order=order,
                    product=product,
                    sku=product.sku,
                    name=product.name,
                    price=price,
                    quantity=quantity,
                )
            order.recalculate_total()
            result.created_orders.append(order.pk)

        if dry_run:
            transaction.set_rollback(True)

    return result
