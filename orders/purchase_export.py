"""Build a supplier replenishment order as an .xlsx workbook."""

from io import BytesIO

import openpyxl
from django.utils import timezone
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("№", lambda i, it: i),
    ("Код 1С", lambda i, it: it.sku),
    ("Артикул", lambda i, it: it.product.article if it.product_id else ""),
    ("Наименование", lambda i, it: it.name),
    ("Заказать", lambda i, it: it.quantity),
    ("Остаток", lambda i, it: it.current_stock),
    ("Продаж/день", lambda i, it: float(it.avg_daily_sales)),
    ("Рекомендовано", lambda i, it: it.suggested_qty),
]


def build_purchase_order_xlsx(order):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказ"

    ws["A1"] = f"Заказ поставщику №{order.pk} — {order.supplier}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Дата: {timezone.localtime(order.created_at):%d.%m.%Y %H:%M}   "
        f"Срок поставки: {order.delivery_days} дн.   "
        f"Страховой запас: {order.safety_days} дн.   "
        f"Период продаж: {order.sales_period_days} дн."
    )

    header_row = 4
    for col, (title, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True)

    items = list(order.items.select_related("product"))
    for i, item in enumerate(items, start=1):
        for col, (_, getter) in enumerate(COLUMNS, start=1):
            ws.cell(row=header_row + i, column=col, value=getter(i, item))

    total_row = header_row + len(items) + 1
    ws.cell(row=total_row, column=4, value="Итого:").font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=order.total_qty).font = Font(bold=True)

    widths = [5, 14, 16, 48, 10, 10, 12, 14]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build_supply_xlsx(rows, stock_warehouses, supplier, period):
    """Export the full replenishment table (all columns, including the
    per-warehouse stock and sales breakdowns) as .xlsx bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказ поставщику"

    ws["A1"] = f"Заказ поставщику — {supplier}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Сформировано: {timezone.localtime():%d.%m.%Y %H:%M}   "
        f"Период продаж: {period} дн."
    )

    headers = ["Бренд", "Товар", "Артикул", "Код 1С", "Остаток"]
    headers += ["Остаток " + w["label"] for w in stock_warehouses]
    headers += [
        f"Продажи за {period} дн.",
        "Продажи Нск",
        "Продажи Кем",
        "Продажи Нкз",
        "Продажи Ozon",
        "Продаж/день",
        "Упаковка",
        "Рекомендовано",
        "Предупреждение",
    ]

    header_row = 4
    for col, title in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=title).font = Font(bold=True)

    for i, r in enumerate(rows, start=1):
        values = [r["brand"], r["name"], r["article"], r["sku"], r["stock"]]
        values += [s["val"] for s in r["wh_stock"]]
        values += [
            r["sold"],
            r["nsk"],
            r["kem"],
            r["nkz"],
            r["ozon"],
            float(r["ads"]),
            r["pack_quantity"],
            r["suggested"],
            r["warning"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=header_row + i, column=col, value=value)

    # Widths: wider for text columns, narrow for the numeric ones.
    text_cols = {1, 2, 3, 4, len(headers)}  # brand, name, article, sku, warning
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            30 if col in text_cols else 11
        )
    ws.column_dimensions["B"].width = 36

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
