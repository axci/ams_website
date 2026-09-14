"""Export the (filtered) sales statistics to an .xlsx workbook."""

from decimal import Decimal
from io import BytesIO

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from django.utils import timezone

# weight_unit variants → factor to kilograms (kept in sync with orders.views).
_GRAM = {"г", "гр", "г.", "гр.", "грамм", "граммов", "Г", "Гр", "ГР", "g", "gr", "gram"}
_TONNE = {"т", "т.", "тн", "тонна", "тонн", "Т", "t", "T", "ton"}


def _kg_factor(unit):
    u = (unit or "").strip()
    if u in _GRAM:
        return Decimal("0.001")
    if u in _TONNE:
        return Decimal("1000")
    return Decimal("1")


def _weight_kg(r):
    if r.product_id and r.product.weight is not None:
        return float(r.quantity * r.product.weight * _kg_factor(r.product.weight_unit))
    return ""


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


COLUMNS = [
    ("Дата", lambda r: timezone.localtime(r.date).strftime("%d.%m.%Y %H:%M") if r.date else ""),
    ("Склад", lambda r: r.warehouse.name),
    ("Код 1С", lambda r: r.sku),
    ("Товар", lambda r: r.product.name if r.product_id else ""),
    ("Бренд", lambda r: r.product.brand.name if r.product_id and r.product.brand_id else ""),
    ("Категория", lambda r: r.product.category.name if r.product_id and r.product.category_id else ""),
    ("Подкатегория", lambda r: r.product.subcategory.name if r.product_id and r.product.subcategory_id else ""),
    ("Модель", lambda r: r.product.model_product.name if r.product_id and r.product.model_product_id else ""),
    ("Поставщик", lambda r: r.product.supplier if r.product_id else ""),
    ("Контрагент", lambda r: r.client),
    ("Тип контрагента", lambda r: r.client_type),
    ("Документ", lambda r: r.document),
    ("Тип документа", lambda r: r.document_type),
    ("Количество", lambda r: float(r.quantity)),
    ("Вес, кг", _weight_kg),
]


def build_sales_xlsx(records):
    """Return the given sales records as .xlsx bytes, one row per record."""
    records = records.select_related(
        "warehouse",
        "product",
        "product__brand",
        "product__category",
        "product__subcategory",
        "product__model_product",
    ).order_by("-date")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.append([name for name, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for r in records.iterator():
        ws.append([_clean(get(r)) for _, get in COLUMNS])

    for idx, (name, _) in enumerate(COLUMNS, start=1):
        width = 40 if name in ("Товар", "Документ") else max(12, min(len(name) + 4, 24))
        ws.column_dimensions[get_column_letter(idx)].width = width

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
