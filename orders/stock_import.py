"""Reconstruct day-by-day stock history from the 1C «Ведомость по товарам».

For each (product, warehouse) the report gives the period's opening balance
(нач. остаток, at the warehouse level) and every stock movement as a document
whose кон. остаток is the running end-of-day balance. We walk each day of the
period (earliest → latest document date across all products) carrying the
balance forward, updating it on days that have documents. Pairs with no
documents keep their opening balance for the whole period.

Only the three mapped warehouses are tracked; transit warehouses are ignored.
"""

from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from catalog.models import Product
from warehouses.models import Warehouse

from .models import StockSnapshot

# Excel warehouse name → site warehouse name.
WAREHOUSE_MAP = {
    "Оптовый Кемерово (АМС)": "Кемерово",
    "Новокузнецк (АМС)": "Новокузнецк",
    "Склад Новосибирск (АМС)": "Новосибирск",
}

COL_A, COL_DATE, COL_SKU, COL_NACH, COL_KON = 1, 2, 4, 5, 8


@dataclass
class StockImportResult:
    snapshots: int = 0
    pairs: int = 0
    period_start: object = None
    period_end: object = None
    unmatched_sku: int = 0
    errors: list = field(default_factory=list)


def _num(value):
    if value is None or str(value).strip() == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ".").strip())
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if value is None or str(value).strip() == "":
        return None
    s = str(value).strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def import_stock(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active
    result = StockImportResult()

    # Parse the hierarchy into { (sku, wh_name): {"nach": Decimal, "docs": [(date, kon)]} }.
    pairs = {}
    all_dates = []
    cur_sku = cur_wh = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, COL_A).value
        if a is None or str(a).strip() == "":
            continue
        a = str(a).strip()
        level = ws.row_dimensions[r].outline_level if r in ws.row_dimensions else 0
        if level == 0:
            v = ws.cell(r, COL_SKU).value
            cur_sku = str(v).strip() if v is not None else None
            cur_wh = None
        elif level == 1:
            cur_wh = WAREHOUSE_MAP.get(a)  # None for transit / unmapped
            if cur_wh and cur_sku:
                pairs[(cur_sku, cur_wh)] = {
                    "nach": _num(ws.cell(r, COL_NACH).value),
                    "docs": [],
                }
        elif level == 2:
            key = (cur_sku, cur_wh)
            if cur_wh and key in pairs:
                d = _date(ws.cell(r, COL_DATE).value)
                if d is not None:
                    pairs[key]["docs"].append((d, _num(ws.cell(r, COL_KON).value)))
                    all_dates.append(d)

    if not all_dates:
        result.errors.append((0, "В файле не найдено документов с датами."))
        return result

    period_start, period_end = min(all_dates), max(all_dates)
    result.period_start, result.period_end = period_start, period_end
    result.pairs = len(pairs)
    day_count = (period_end - period_start).days + 1

    site_wh = {w.name: w for w in Warehouse.objects.all()}
    target_wh = {name: site_wh.get(name) for name in set(WAREHOUSE_MAP.values())}
    products = {p.sku: p for p in Product.objects.all()}

    snapshots = []
    for (sku, wh_name), info in pairs.items():
        warehouse = target_wh.get(wh_name)
        if warehouse is None:
            continue
        product = products.get(sku)
        if product is None:
            result.unmatched_sku += 1
        # End-of-day balance for each date that has documents (last one wins).
        eod = {}
        for d, kon in sorted(info["docs"], key=lambda x: x[0]):
            eod[d] = kon
        balance = info["nach"]
        for i in range(day_count):
            day = period_start + timedelta(days=i)
            if day in eod:
                balance = eod[day]
            snapshots.append(
                StockSnapshot(
                    sku=sku, product=product, warehouse=warehouse,
                    date=day, quantity=balance,
                )
            )

    warehouses = [w for w in target_wh.values() if w is not None]
    with transaction.atomic():
        StockSnapshot.objects.filter(
            warehouse__in=warehouses, date__gte=period_start, date__lte=period_end
        ).delete()
        StockSnapshot.objects.bulk_create(snapshots, batch_size=2000)
    result.snapshots = len(snapshots)
    return result
