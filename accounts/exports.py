"""Export users to an .xlsx workbook (admin-only), with all their info."""

from io import BytesIO

import openpyxl
from django.utils import timezone
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _moment(value):
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M") if value else ""


def _companies(user):
    return list(user.companies.all())


# Header, getter. One row per user; M2M values are joined into one cell.
COLUMNS = [
    ("Логин", lambda u: u.username),
    ("Email", lambda u: u.email),
    ("Имя", lambda u: u.first_name),
    ("Фамилия", lambda u: u.last_name),
    ("Активен", lambda u: u.is_active),
    ("Персонал", lambda u: u.is_staff),
    ("Суперпользователь", lambda u: u.is_superuser),
    ("Менеджер", lambda u: str(u.manager) if u.manager_id else ""),
    ("Тип цены", lambda u: u.price_type.name if u.price_type_id else ""),
    ("Организация 1С", lambda u: u.organization_1c),
    ("Точный остаток", lambda u: u.show_stock),
    ("Доступ: статистика продаж", lambda u: u.can_view_sales),
    ("Доступ: заказ поставщику", lambda u: u.can_order_supply),
    ("Мин. сумма доставки", lambda u: u.free_delivery_min),
    ("Склады", lambda u: ", ".join(w.name for w in u.warehouses.all())),
    ("Компании", lambda u: "; ".join(
        c.company_name or c.code or f"#{c.pk}" for c in _companies(u)
    )),
    ("Коды 1С компаний", lambda u: ", ".join(c.code for c in _companies(u) if c.code)),
    ("ИНН компаний", lambda u: ", ".join(c.inn for c in _companies(u) if c.inn)),
    ("Задолженность", lambda u: u.total_debt()),
    ("Дата регистрации", lambda u: _moment(u.date_joined)),
    ("Последний вход", lambda u: _moment(u.last_login)),
]


def build_users_xlsx(users):
    """Return the given users as .xlsx bytes (one row each, all columns)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Пользователи"

    headers = [title for title, _ in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for user in users:
        ws.append([get(user) for _, get in COLUMNS])

    for idx, title in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(
            12, min(len(title) + 4, 32)
        )

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
