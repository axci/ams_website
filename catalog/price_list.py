"""Build and email a price list (.xlsx) for a price type + warehouse.

Columns: Артикул, Бренд, Наименование, Свободный остаток (stock on the chosen
warehouse) and Цена (the chosen price type's price, falling back to the
product's base price when that type has no explicit price). Only active
products are included.
"""

from io import BytesIO

import openpyxl
from django.conf import settings
from django.core.mail import EmailMessage
from django.db.models import DecimalField, F, IntegerField, OuterRef, Subquery, Value
from django.db.models.functions import Coalesce
from django.utils import timezone
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from warehouses.models import Stock

from .models import Product, ProductPrice

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

HEADERS = ["Артикул", "Бренд", "Наименование", "Свободный остаток", "Цена"]

# Brand names shown differently in the exported price list (display only — the
# catalog keeps the original name).
BRAND_LABEL_OVERRIDES = {
    "SCT (Germany)": "SCT",
    "Аляska": "Alaska",
}


def _price_list_products(price_type, warehouse, in_stock_only=False):
    """Active products annotated with the type's price and the warehouse stock."""
    price_sub = ProductPrice.objects.filter(
        product=OuterRef("pk"), price_type=price_type
    ).values("price")[:1]
    stock_sub = Stock.objects.filter(
        product=OuterRef("pk"), warehouse=warehouse
    ).values("quantity")[:1]
    products = (
        Product.objects.filter(is_active=True)
        .select_related("brand")
        .annotate(
            list_price=Coalesce(
                Subquery(
                    price_sub,
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                ),
                F("price"),
            ),
            free_stock=Coalesce(
                Subquery(stock_sub, output_field=IntegerField()), Value(0)
            ),
        )
    )
    # Never list a product with no price.
    products = products.filter(list_price__gt=0)
    if in_stock_only:
        products = products.filter(free_stock__gt=0)
    return products.order_by("brand__name", "name")


def build_price_list_xlsx(price_type, warehouse, in_stock_only=False):
    """Return the price list as .xlsx bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    today = timezone.localdate().strftime("%d.%m.%Y")

    ws.append([f"Прайс «{price_type.name}» · склад {warehouse.name} · {today}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(row=1, column=1)
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.append(HEADERS)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A3"

    for p in _price_list_products(price_type, warehouse, in_stock_only).iterator():
        brand = p.brand.name if p.brand_id else ""
        ws.append([
            p.article,
            BRAND_LABEL_OVERRIDES.get(brand, brand),
            p.name,
            p.free_stock,
            float(p.list_price or 0),
        ])

    for i, width in enumerate([18, 22, 50, 18, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def send_price_list_email(price_type, warehouse, recipients, in_stock_only=False):
    """Build the price list and email it as an .xlsx attachment.

    `recipients` is a list of addresses. Raises on send failure."""
    data = build_price_list_xlsx(price_type, warehouse, in_stock_only)
    filename = f"price-list-{timezone.localdate().strftime('%Y-%m-%d')}.xlsx"
    subject = f"Прайс-лист «{price_type.name}» — Автомеханика-Сибирь"
    body = (
        f"Во вложении — прайс-лист «{price_type.name}».\n"
        f"Склад: {warehouse.name}.\n"
        f"Дата формирования: {timezone.localdate().strftime('%d.%m.%Y')}.\n"
    )
    msg = EmailMessage(
        subject=subject,
        body=body,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=list(recipients),
    )
    msg.attach(filename, data, XLSX_CONTENT_TYPE)
    msg.send()
